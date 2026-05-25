from pathlib import Path

import pdfplumber
from docx import Document as DocxDocument


def extract_text_from_file(path: Path, filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".txt") or lower.endswith(".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    if lower.endswith(".docx"):
        doc = DocxDocument(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if lower.endswith(".pdf"):
        parts: list[str] = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(text)
        return "\n\n".join(parts)
    if lower.endswith(".xlsx"):
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        rows: list[str] = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append(" | ".join(cells))
        return "\n".join(rows)
    raise ValueError(f"Unsupported file type: {filename}")
